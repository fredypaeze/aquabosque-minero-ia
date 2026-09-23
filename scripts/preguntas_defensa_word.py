"""Preguntas de defensa técnica (Q&A) -> Word."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches

VERDE = "1B5E20"; NAVY = "123A5C"; GRIS = "5B6B62"; INK = "222A26"; AMBAR = "B45309"

QA = [
 ("¿A qué se refiere el índice de “sensibilidad”?",
  "Es el valor ambiental y social que está en juego en el territorio: cuánto hay que perder si algo sale mal. "
  "Se construye con las hectáreas de áreas protegidas del RUNAP cercanas al municipio (normalizadas), más un "
  "incremento si el municipio es PDET (posconflicto). Va de 0 a 1: 0 = poco valor a proteger, 1 = mucho. "
  "No mide daño; mide la importancia de lo que está expuesto.",
  "Frase corta: “Sensibilidad es cuánto valor ambiental y social hay que proteger allí. Un territorio con "
  "parques naturales o en posconflicto es más sensible: una presión ahí pesa más.”"),
 ("¿De qué fuente viene la señal de fuego?",
  "De la NASA, sistema FIRMS (Fire Information for Resource Management System). En concreto, de los sensores "
  "satelitales VIIRS (375 m, en los satélites Suomi-NPP y NOAA-20) y MODIS (1 km, en Terra/Aqua). Detecta focos de "
  "calor activos casi en tiempo real; nosotros los contamos por municipio en los últimos 7 días y sumamos su "
  "intensidad (FRP, en megavatios). Es dato abierto y se actualiza a diario.",
  "Frase corta: “El fuego viene del satélite de la NASA (FIRMS); son focos de calor activos, actualizados a diario.”"),
 ("¿Cuál es la arquitectura y el software de procesamiento?",
  "Todo en Python de código abierto. Integración territorial sin GDAL (cruce por código DANE exacto o por centroide "
  "con distancia haversine). Modelo XGBoost multiclase con explicabilidad SHAP; capas de rigor con Conformal "
  "Prediction e Isolation Forest. Satélite profundo: Sentinel-2 vía STAC + rasterio y un U-Net en PyTorch sobre GPU "
  "NVIDIA L40S. Asistente con LLM local (Llama 3.3 70B / Qwen 2.5) y embeddings bge-m3. Aplicación en Streamlit + "
  "Plotly. Repositorio abierto en GitHub con CI y 16 pruebas. Todo corre en la infraestructura del Ministerio.",
  "Frase corta: “Python abierto, XGBoost + SHAP para el modelo, PyTorch en GPU para el satélite y un LLM local "
  "para el asistente. Corre en la infraestructura del Ministerio: el dato no sale del Estado.”"),
 ("¿Cómo se leen los índices? ¿0 es malo y 1 es bueno?",
  "Es al revés de lo intuitivo: todos los índices van de 0 a 1, donde 1 = mayor alerta y prioridad de revisión, y "
  "0 = sin señal. No es “0 malo, 1 bueno”. idx_minero (presión minera), idx_deforestacion (pérdida de bosque), "
  "idx_fuego (quema reciente), idx_hidrico (agua degradada = 1 − ICA) e idx_sensibilidad (valor a proteger) apuntan "
  "todos en la misma dirección: más cerca de 1, más urgente.",
  "Matiz honesto clave: en idx_hidrico, un 0 puede significar “sin estación de medición” (sin dato observado), "
  "NO “agua sana”. El sistema lo marca como ausencia de dato; no inventa que el agua está bien."),
]

NOTA_IA = ("El asistente de IA de la aplicación (RAG con LLM local) también responde estas preguntas en vivo, "
           "aterrizado en la metodología del sistema. Se probó: responde correctamente la fuente del fuego, la "
           "lectura de índices y la definición de sensibilidad.")


def build():
    doc = Document()
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(doc.sections[0], m, Inches(0.9))
    h = doc.add_heading("", level=0)
    r = h.add_run("Preguntas de defensa técnica · AquaBosque Minero IA")
    r.font.size = Pt(19); r.font.color.rgb = RGBColor.from_string(VERDE); r.font.name = "Calibri"
    p = doc.add_paragraph()
    rp = p.add_run("Respuestas respaldadas por el código real del sistema. Úsalas como referencia rápida ante el jurado.")
    rp.font.size = Pt(10.5); rp.font.italic = True; rp.font.color.rgb = RGBColor.from_string(GRIS)
    doc.add_paragraph()

    for i, (q, a, extra) in enumerate(QA, 1):
        hp = doc.add_paragraph()
        rn = hp.add_run(f"{i}.  {q}")
        rn.font.size = Pt(13.5); rn.font.bold = True; rn.font.color.rgb = RGBColor.from_string(NAVY); rn.font.name = "Calibri"
        ap = doc.add_paragraph()
        ra = ap.add_run(a); ra.font.size = Pt(12); ra.font.color.rgb = RGBColor.from_string(INK); ra.font.name = "Calibri"
        ap.paragraph_format.line_spacing = 1.2
        ep = doc.add_paragraph()
        re = ep.add_run(extra); re.font.size = Pt(11.5); re.font.italic = True; re.font.color.rgb = RGBColor.from_string(AMBAR); re.font.name = "Calibri"
        ep.paragraph_format.space_after = Pt(14)

    doc.add_paragraph()
    np_ = doc.add_paragraph()
    rnn = np_.add_run("💡 " + NOTA_IA)
    rnn.font.size = Pt(11); rnn.font.color.rgb = RGBColor.from_string(VERDE); rnn.font.name = "Calibri"

    out = Path(__file__).resolve().parents[1] / "outputs" / "jurado_2026" / "PREGUNTAS_DEFENSA_TECNICA.docx"
    doc.save(out); print("OK:", out)


if __name__ == "__main__":
    build()
