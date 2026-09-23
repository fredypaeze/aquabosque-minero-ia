"""Guion de presentación actualizado (con lámina Rigor Avanzado) -> Word."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

VERDE = "1B5E20"; NAVY = "123A5C"; GRIS = "5B6B62"; INK = "222A26"; AMBAR = "B45309"

SEG = [
 ("1", "Portada — AquaBosque Minero IA",
  "Bienvenidos. Hoy presentamos AquaBosque Minero IA. En pocas palabras, es una herramienta inteligente que nos dice qué municipios de Colombia necesitan atención urgente por riesgos ambientales. Lo hace cruzando información de minería, deforestación y calidad del agua en una sola plataforma, usando únicamente datos oficiales."),
 ("2", "El Problema",
  "El problema hoy no es que nos falte información. Colombia tiene muchísimos datos ambientales, pero están dispersos. Revisar por un lado la minería, por otro los bosques y por otro el agua, hace muy difícil saber dónde actuar primero. No venimos a proponer otro tablero aislado más, sino una lectura integrada del territorio para que el Ministerio pueda tomar decisiones rápidas."),
 ("3", "La Propuesta",
  "Nuestra solución ya es un producto funcional, no un concepto. Tomamos esos datos dispersos, los unimos a nivel municipal y una Inteligencia Artificial les asigna un nivel de riesgo. Lo más valioso de esta herramienta es que hace que la priorización de los territorios sea operativa, transparente y muy fácil de entender para cualquier analista o directivo."),
 ("4", "Datos Abiertos Integrados",
  "Para que esto tenga validez institucional, la regla de oro fue usar solo fuentes oficiales y verificables. Integramos bases de la Agencia Nacional de Minería, el IDEAM y registros de áreas protegidas. Aquí no hay estimaciones sin fundamento ni datos inventados; todo es auditable y se puede rastrear hasta su origen."),
 ("5", "Construcción de la Unidad Territorial",
  "¿Cómo mezclamos peras con manzanas? Estandarizamos todo al nivel de los 1.122 municipios del país. Cruzamos los datos de minería y agua usando las coordenadas de cada lugar. Si en algún municipio no hay datos de calidad de agua, el sistema es honesto y lo marca como ausencia de dato, no inventa problemas donde no hay evidencia."),
 ("6", "Cómo Funciona la IA",
  "El motor detrás de esto clasifica el riesgo en cuatro niveles: Bajo, Medio, Alto y Crítico. Pero la verdadera magia de este modelo no es solo decirnos el nivel de riesgo, sino explicarnos el por qué. El sistema nos muestra exactamente qué factores —si fue el fuego, la minería o la deforestación— pesaron más para encender la alerta en un municipio específico."),
 ("7", "Priorización e Interpretabilidad (Caso Barrancabermeja)",
  "Veamos un caso real: Barrancabermeja. El sistema lo marca en riesgo 'Crítico'. Al mirar el detalle, vemos que tiene actividad minera, afectación hídrica y alta sensibilidad ambiental al mismo tiempo. La IA no reemplaza al experto ni sanciona automáticamente; lo que hace es decirle a la autoridad: 'Revisen aquí primero, esta es la evidencia combinada'."),
 ("8", "Rigor Avanzado (Confianza calibrada)  ★ NUEVA",
  "Y aquí está lo que de verdad nos diferencia. Cualquier modelo le entrega una 'confianza', pero eso es solo la probabilidad interna del sistema. Nosotros vamos un paso más allá: con una técnica llamada conformal prediction, le damos una garantía estadística —el nivel de riesgo real de un municipio cae dentro de nuestra predicción en el 90% de los casos, y lo verificamos: cumplió el 91.4%. Y como segundo control, un modelo distinto, de forma independiente y sin mirar nuestras etiquetas, confirma 43 de los municipios que priorizamos y descubre 14 más, atípicos, que la fórmula por sí sola no resaltaba. Este es el tipo de rigor que se usa en medicina y en finanzas y que rara vez se ve en el sector público. No solo priorizamos: sabemos con qué certeza lo hacemos."),
 ("9", "Capa Satelital · Monitoreo Near-Real-Time",
  "Y no nos quedamos solo con datos históricos. Le conectamos una señal de satélite de la NASA que detecta focos de calor casi en tiempo real. Si el sistema dice que un municipio es de alto riesgo, y hoy mismo el satélite detecta fuego allí, cruzamos esa información al instante. Es pasar del análisis estático a la respuesta en tiempo real."),
 ("10", "Capa Satelital · Deforestación (GPU)",
  "También tenemos la capacidad de hacer un zoom satelital profundo usando imágenes procesadas en la infraestructura del Ministerio. Esto nos permite ver el antes y el después de una zona y calcular exactamente cuántas hectáreas de bosque se acaban de perder. Todo el procesamiento de las imágenes se hace en casa, garantizando la soberanía de nuestros datos."),
 ("11", "IA Generativa Soberana (Asistente)",
  "Para facilitar aún más las cosas, integramos un asistente virtual. Usted puede preguntarle en lenguaje natural, por ejemplo: '¿Por qué La Macarena está priorizada?'. El asistente lee los datos del sistema y le explica la situación. Está configurado para no inventar nada; solo responde basado en la evidencia que ya tenemos recolectada."),
 ("12", "La Aplicación Funcional (Mapa y Ranking)",
  "Todo esto vive en una aplicación real y fácil de usar. Tenemos un mapa nacional de calor, un ranking con el 'Top 15' de municipios críticos y fichas detalladas por territorio. La navegación está diseñada para que en menos de dos minutos usted encuentre la información que necesita para justificar una acción."),
 ("13", "Demostración",
  "Hagamos un recorrido rápido: entramos, vemos el mapa nacional, pasamos al ranking, abrimos un municipio, revisamos la explicación de por qué está priorizado y descargamos el resultado. En menos de dos minutos se ve el valor completo. (Si la conexión falla, tenemos las capturas de respaldo)."),
 ("14", "Caso de Uso Institucional (Descarga)",
  "Sabemos que los datos deben fluir. Por eso, cualquier consulta o ranking se puede descargar directamente en un archivo de Excel o CSV con un solo clic. El objetivo es que los equipos técnicos puedan llevarse esta información y cruzarla rápidamente con sus propios reportes internos."),
 ("15", "Qué Valida el MVP",
  "Esta versión ya fue auditada técnicamente y funciona: datos reales integrados, un modelo que se puede auditar, el satélite operando, la aplicación desplegada y todo el código abierto en un repositorio público. Y lo que todavía falta, lo decimos con claridad."),
 ("16", "Escalamiento (Ruta)",
  "Nuestro siguiente paso lógico es automatizar que los datos se actualicen solos y generar alertas directas a la autoridad. Y sabemos exactamente cuál es la frontera hacia la que vamos: modelos espaciales bayesianos, detección de deforestación por radar como la de Brasil, e inferencia causal. No prometemos una infraestructura gigantesca desde el día uno, sino un crecimiento comprobable, sobre una base ya validada."),
 ("17", "Impacto Potencial",
  "El impacto es concreto: reducimos drásticamente las horas de revisión manual y enfocamos los recursos del Estado donde más se necesitan, con evidencia clara y unificada. Y es una base escalable para priorizar regiones sensibles como la Amazonía, la Orinoquía y el Pacífico."),
 ("18", "Cierre",
  "En conclusión, AquaBosque Minero IA convierte datos dispersos en decisiones estratégicas. Es un producto real, con inteligencia artificial de satélite y generativa, corriendo en infraestructura pública y con el dato dentro del Estado. Muchas gracias; quedamos atentos a sus preguntas."),
]


def build():
    doc = Document()
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(doc.sections[0], m, Inches(0.9))
    h = doc.add_heading("", level=0)
    r = h.add_run("Guion de presentación · AquaBosque Minero IA")
    r.font.size = Pt(20); r.font.color.rgb = RGBColor.from_string(VERDE); r.font.name = "Calibri"
    p = doc.add_paragraph()
    rp = p.add_run("Concurso Datos al Ecosistema 2026 · alineado a las 18 láminas del deck (incluye la nueva lámina 8, Rigor Avanzado).")
    rp.font.size = Pt(10.5); rp.font.italic = True; rp.font.color.rgb = RGBColor.from_string(GRIS)
    doc.add_paragraph()

    for num, titulo, texto in SEG:
        hp = doc.add_paragraph()
        rn = hp.add_run(f"Diapositiva {num} · {titulo}")
        rn.font.size = Pt(13); rn.font.bold = True
        rn.font.color.rgb = RGBColor.from_string(AMBAR if "NUEVA" in titulo else NAVY)
        rn.font.name = "Calibri"
        tp = doc.add_paragraph()
        rt = tp.add_run("“" + texto + "”")
        rt.font.size = Pt(12.5); rt.font.color.rgb = RGBColor.from_string(INK); rt.font.name = "Calibri"
        tp.paragraph_format.space_after = Pt(12); tp.paragraph_format.line_spacing = 1.25

    out = Path(__file__).resolve().parents[1] / "outputs" / "jurado_2026" / "GUION_ACTUALIZADO.docx"
    doc.save(out); print("OK:", out)


if __name__ == "__main__":
    build()
